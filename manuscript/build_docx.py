#!/usr/bin/env python
"""Assemble AdaSpark_CaseStudy_Report.docx inside the official course template.

The official template (manuscript/template/official_template.docx, downloaded from the link in
the manuscript guidelines) is a course cover page plus Index / List of Figures / List of Tables /
Contributions pages, followed by an IEEE-Access-style paper: title and author block, ABSTRACT,
INDEX TERMS, then a two-column body. This builder keeps the cover and front matter, fills the
front-matter tables in the template's own cell formatting, replaces the template's sample paper
with sections/*.md, and formats every element by cloning the template's own paragraphs (headings,
body text, equations, captions). Full-width figures and tables are set in single-column blocks
between continuous section breaks, as IEEE layouts do.

Front matter: ABSTRACT is sections/00-abstract.md (one paragraph, 150-250 words); the Executive
Summary (00-exec-summary.md, with Table 1) follows INDEX TERMS. Fields the supplied files cannot
fill (student names, registration numbers, e-mails, Team ID, per-member contributions) are written
as highlighted [ ... ] flags for manual completion; nothing is invented.

Page numbers: the Index / List of Figures / List of Tables cells are PAGEREF fields. When Microsoft
Word is available the builder lets Word update and paginate a temporary copy, reads the computed
page numbers back and writes them into the fields' cached results, so the delivered file shows
real page numbers without an update prompt and without being re-saved by Word (which can
recompress the 600 DPI figures). Without Word the fields keep "-" and are flagged for update.

Equations: editable Microsoft Equation Editor (OMML) objects converted from equations.json with
Office's MML2OMML.XSL. Requires python-docx and lxml (manuscript toolchain interpreter).
The build fails loudly if any figure, table, caption, equation, reference or word-count check fails.
"""
import copy
import json
import os
import re
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import etree


def n_words(text):
    """Words as rendered: markdown table pipes and |---| rows never reach the document."""
    return sum(1 for w in text.split() if re.search(r"\w", w))


HERE = Path(__file__).resolve().parent
SECTIONS = HERE / "sections"
DIAGRAMS = HERE / "diagrams"
TEMPLATE = HERE / "template" / "official_template.docx"
OUT = HERE / "AdaSpark_CaseStudy_Report.docx"
XSL = Path(os.environ.get("MML2OMML_XSL",
                          r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"))
TITLE = "Self-Adaptive Big Data Programming Using Reinforcement Learning and Spark"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
WORD_BAND = (10_000, 15_000)          # planning table: total words, references excluded

BODY = ["02-introduction.md", "03-problem.md", "04-background.md", "05-requirements.md",
        "06-architecture.md", "07-execution-layer.md", "08-rl-layer.md", "09-scheduling.md",
        "10-self-tuning.md", "11-explainability.md", "12-monitoring.md", "13-governance.md",
        "14-exp-setup.md", "15-metrics.md", "16-results.md", "17-comparative.md", "18-trust.md",
        "19-usecases.md", "20-business.md", "21-limitations.md", "22-future.md",
        "23-conclusion.md", "../data_availability.md", "24-references.md", "25-appendix.md"]
FIG_AFTER = {1: "02", 2: "02", 3: "03", 4: "04", 5: "05", 6: "06", 7: "06", 8: "07",
             9: "08", 10: "08", 11: "09", 12: "09", 13: "10", 14: "10", 15: "11",
             16: "11", 17: "12", 18: "13", 19: "14", 20: "15", 21: "16", 22: "16",
             23: "16", 24: "17", 25: "18", 26: "19", 27: "19", 28: "20", 29: "21",
             30: "22", 31: "22", 32: "23"}
# template paragraph indices of the prototypes (verified against the downloaded template)
P_COVER_TITLE, P_TITLE, P_AUTHORS, P_EMAILS = 1, 89, 90, 92
P_ABSTRACT, P_TERMS, P_SECTION0_END = 95, 96, 98
P_H1, P_BODY, P_BODY2, P_H2, P_EQ, P_FIGCAP, P_TABCAP = 99, 100, 101, 107, 125, 133, 149
# the template pushes each front-matter part onto its own page with runs of empty paragraphs,
# which only works for its own table lengths; they are replaced by page-break-before on the
# part's heading (List of Figures, List of Tables, Contributions, paper title)
SPACERS = [*range(21, 28), *range(30, 50), *range(52, 71), *range(74, 89)]
PAGE_STARTS = (28, 50, 71, 89)
COVER_TITLE_SZ = "44"   # half-points; the template's 32 pt placeholder wraps the full title onto
                        # four lines and pushes the cover's logo block onto a second page
STUDENTS = 4            # template author slots 2-5 and the consolidated table's four rows
FLAG_RE = re.compile(r"(\[[^\]]*\b(?:to be (?:inserted|added|entered))\b[^\]]*\])")

EQUATION_LINE = re.compile(r"^(?P<body>.*\S)\s+\((?P<n>\d{1,2})\)\s*$")
TABLE_CAPTION = re.compile(r"^##\s+Table\s+(\d+):\s*(.+)$")
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|(?<![\w*])\*[^*\s][^*]*\*(?![\w*]))")
ROMAN = ["", "I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII", "XIII",
         "XIV", "XV", "XVI", "XVII", "XVIII", "XIX", "XX", "XXI", "XXII"]
# OOXML child order (ECMA-376 CT_PPr / CT_RPr); Word rejects out-of-order properties.
PPR_SEQ = ("pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr "
           "suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct "
           "topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid spacing ind "
           "contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment "
           "textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange").split()
RPR_SEQ = ("rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint "
           "noProof snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u "
           "effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath").split()
RPR_ALLOWED = set(RPR_SEQ)

PS_WORD_UPDATE = r"""
$ErrorActionPreference = 'Stop'
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$word.DisplayAlerts = 0
try {
  $doc = $word.Documents.Open('%(src)s', $false, $false)
  $doc.Fields.Update() | Out-Null
  $doc.Repaginate()
  $doc.Fields.Update() | Out-Null
  $pages = $doc.ComputeStatistics(2)
  $words = $doc.ComputeStatistics(0)
  $doc.SaveAs2('%(dst)s', 16)
  $doc.Close([ref]0)
  Write-Output "STATS $pages $words"
} finally { $word.Quit() }
"""


def ordered_insert(parent, el, seq):
    """Replace any same-tag child, then insert el at its schema position."""
    tag = el.tag.split("}")[1]
    for old in parent.findall(el.tag):
        parent.remove(old)
    rank = seq.index(tag)
    for child in parent:
        ctag = child.tag.split("}")[1]
        if ctag in seq and seq.index(ctag) > rank:
            child.addprevious(el)
            return el
    parent.append(el)
    return el


def highlight(rpr):
    el = OxmlElement("w:highlight")
    el.set(qn("w:val"), "yellow")
    return ordered_insert(rpr, el, RPR_SEQ)


def figure_refs(text):
    """Template rule: in-text figure references use the abbreviation "Fig."."""
    text = re.sub(r"\bFigures (?=\d)", "Figs. ", text)
    return re.sub(r"\bFigure (?=\d)", "Fig. ", text)


def read_section(fn):
    return (SECTIONS / fn).resolve().read_text(encoding="utf-8")


class Builder:
    def __init__(self):
        if not XSL.exists():
            sys.exit(f"MML2OMML.XSL not found at {XSL}; set MML2OMML_XSL")
        self.doc = Document(str(TEMPLATE))
        paras = self.doc.paragraphs
        self.proto = {k: copy.deepcopy(paras[i]._p) for k, i in
                      (("h1", P_H1), ("body", P_BODY), ("body2", P_BODY2), ("h2", P_H2),
                       ("eq", P_EQ), ("figcap", P_FIGCAP), ("tabcap", P_TABCAP))}
        self.fm = {i: paras[i]._p for i in (*SPACERS, *PAGE_STARTS, P_COVER_TITLE, P_TITLE,
                                            P_AUTHORS, P_EMAILS, P_ABSTRACT, P_TERMS,
                                            P_TERMS + 1)}
        self.xslt = etree.XSLT(etree.parse(str(XSL)))
        self.mathml = {k: v for k, v in json.loads(
            (HERE / "equations.json").read_text(encoding="utf-8")).items() if k.isdigit()}
        self.captions = {int(k): v for k, v in json.loads(
            (HERE / "figure_captions.json").read_text(encoding="utf-8")).items() if k.isdigit()}
        body = self.doc.element.body
        kids = list(body)
        sect0_end = paras[P_SECTION0_END]._p
        self.col1_sectpr = copy.deepcopy(sect0_end.find(qn("w:pPr")).find(qn("w:sectPr")))
        type_el = self.col1_sectpr.find(qn("w:type"))
        if type_el is None:
            type_el = OxmlElement("w:type")
            self.col1_sectpr.insert(0, type_el)
        type_el.set(qn("w:val"), "continuous")
        two_col_end = [el for el in kids if el.tag == qn("w:p") and el.find(qn("w:pPr")) is not None
                       and el.find(qn("w:pPr")).find(qn("w:sectPr")) is not None]
        self.anchor = two_col_end[1]               # closes the template's two-column section
        self.col2_sectpr = copy.deepcopy(self.anchor.find(qn("w:pPr")).find(qn("w:sectPr")))
        start = kids.index(sect0_end) + 1
        for el in kids[start:]:                    # drop the template's sample paper body
            if el is not self.anchor and el.tag != qn("w:sectPr"):
                body.remove(el)
        self.bookmark_id = 100
        self.equations, self.figures, self.tables, self.sections = [], [], [], []
        self.pagerefs = []                         # (bookmark, cached-result w:t)
        self.in_block = False
        self.first_par = True                      # next body paragraph follows a heading

    # ---------- primitives ----------
    def _clone(self, key, keep_runs=False):
        p = copy.deepcopy(self.proto[key])
        if not keep_runs:
            for child in list(p):
                if child.tag != qn("w:pPr"):
                    p.remove(child)
        return p

    def _run(self, text, rpr_from=None, bold=False, italic=False, code=False, size=None,
             flag=False):
        r = OxmlElement("w:r")
        if rpr_from is not None and rpr_from.find(qn("w:rPr")) is not None:
            rpr = copy.deepcopy(rpr_from.find(qn("w:rPr")))
        else:
            rpr = OxmlElement("w:rPr")
        if bold:
            ordered_insert(rpr, OxmlElement("w:b"), RPR_SEQ)
        if italic:
            ordered_insert(rpr, OxmlElement("w:i"), RPR_SEQ)
        if code:
            fonts = OxmlElement("w:rFonts")
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                fonts.set(qn(a), "Consolas")
            ordered_insert(rpr, fonts, RPR_SEQ)
        if size:
            for tag in ("w:sz", "w:szCs"):
                el = OxmlElement(tag)
                el.set(qn("w:val"), str(size))
                ordered_insert(rpr, el, RPR_SEQ)
        if flag:
            highlight(rpr)
        r.append(rpr)
        for i, part in enumerate(text.split("\t")):
            if i:
                r.append(OxmlElement("w:tab"))
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = part
            r.append(t)
        return r

    def _inline(self, p, text, size=None, base_run=None):
        for tok in INLINE.split(text.replace("\\|", "|")):
            if not tok:
                continue
            if tok.startswith("**") and tok.endswith("**"):
                p.append(self._run(tok[2:-2], base_run, bold=True, size=size))
            elif tok.startswith("`") and tok.endswith("`"):
                p.append(self._run(tok[1:-1], base_run, code=True, size=size))
            elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
                p.append(self._run(tok[1:-1], base_run, italic=True, size=size))
            else:
                for part in FLAG_RE.split(tok):
                    if part:
                        p.append(self._run(part, base_run, size=size, flag=bool(FLAG_RE.fullmatch(part))))

    def _put(self, el):
        self.anchor.addprevious(el)

    def _bookmark(self, p, name):
        start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
        start.set(qn("w:id"), str(self.bookmark_id))
        start.set(qn("w:name"), name)
        end.set(qn("w:id"), str(self.bookmark_id))
        self.bookmark_id += 1
        ppr = p.find(qn("w:pPr"))
        if ppr is not None:
            ppr.addnext(start)
        else:
            p.insert(0, start)
        p.append(end)

    def _section_break(self, sectpr):
        p = OxmlElement("w:p")
        ppr = OxmlElement("w:pPr")
        ppr.append(copy.deepcopy(sectpr))
        p.append(ppr)
        self._put(p)

    def _open_block(self):
        if not self.in_block:
            self._section_break(self.col2_sectpr)      # end the two-column text before it
            self.in_block = True

    def _close_block(self):
        if self.in_block:
            self._section_break(self.col1_sectpr)      # end the single-column block
            self.in_block = False

    # ---------- elements ----------
    def heading(self, level, text, bookmark=None):
        self._close_block()
        p = self._clone("h1" if level == 1 else "h2")
        proto_run = self.proto["h1" if level == 1 else "h2"].find(qn("w:r"))
        p.append(self._run(text, proto_run))
        if bookmark:
            self._bookmark(p, bookmark)
        self._put(p)
        self.first_par = True

    def paragraph(self, text, size=None, hanging=False):
        self._close_block()
        # template: the first paragraph after a heading is flush, later ones take a first-line indent
        p = self._clone("body" if (self.first_par or hanging) else "body2")
        self.first_par = False
        if hanging:
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), "360")
            ind.set(qn("w:hanging"), "360")
            ordered_insert(p.find(qn("w:pPr")), ind, PPR_SEQ)
        self._inline(p, figure_refs(text), size=size, base_run=self.proto["body"].find(qn("w:r")))
        self._put(p)

    def equation(self, line):
        m = EQUATION_LINE.match(line.strip())
        if not m or m["n"] not in self.mathml or not re.search(r"[=←]", m["body"]):
            return False
        self._close_block()
        p = self._clone("eq")
        ppr = p.find(qn("w:pPr"))
        if ppr.find(qn("w:ind")) is not None:
            ppr.remove(ppr.find(qn("w:ind")))
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "right")
        tab.set(qn("w:pos"), "4320")                # right edge of a two-column column
        tabs.append(tab)
        ordered_insert(ppr, tabs, PPR_SEQ)
        omml = self.xslt(etree.fromstring(self.mathml[m["n"]].encode("utf-8"))).getroot()
        p.append(copy.deepcopy(omml))
        p.append(self._run("\t(%s)" % m["n"], self.proto["eq"].find(qn("w:r"))))
        self._put(p)
        self.equations.append(int(m["n"]))
        return True

    def table(self, lines, number=None, title=None):
        self._open_block()
        for text in (("TABLE %d" % number, title) if number else ()):
            cap = self._clone("tabcap")
            cap.append(self._run(text, self.proto["tabcap"].find(qn("w:r"))))
            if text.startswith("TABLE"):
                self._bookmark(cap, "_AdaTab%d" % number)
            ordered_insert(cap.find(qn("w:pPr")), OxmlElement("w:keepNext"), PPR_SEQ)
            self._put(cap)
        rows = [[c.strip() for c in re.split(r"(?<!\\)\|", ln.strip().strip("|"))] for ln in lines]
        if len(rows) > 1 and all(set(c) <= set("-: ") for c in rows[1]):
            del rows[1]
        width = max(len(r) for r in rows)
        t = self.doc.add_table(rows=len(rows), cols=width)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(rows):
            for j in range(width):
                cell = t.cell(i, j)
                par = cell.paragraphs[0]._p
                for child in list(par):
                    if child.tag != qn("w:pPr"):
                        par.remove(child)
                self._inline(par, row[j] if j < len(row) else "", size=16)
                if i == 0:
                    for r in par.findall(qn("w:r")):
                        ordered_insert(r.find(qn("w:rPr")), OxmlElement("w:b"), RPR_SEQ)
                # caption stays with the table; short tables stay whole, long ones may continue
                # on the next page after their first rows, with the header row repeated
                if i < (len(rows) - 1 if len(rows) <= 9 else 2):
                    ordered_insert(par.get_or_add_pPr(), OxmlElement("w:keepNext"), PPR_SEQ)
            trpr = t.rows[i]._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            if i == 0:
                trpr.append(OxmlElement("w:tblHeader"))
        self._put(t._tbl)
        spacer = OxmlElement("w:p")
        self._put(spacer)
        self.first_par = True
        if number:
            self.tables.append(number)

    def figure(self, number, path):
        self._open_block()
        par = self.doc.add_paragraph()
        par.alignment = 1
        par.add_run().add_picture(str(path), width=Inches(6.2))
        par.paragraph_format.keep_with_next = True     # image stays with its caption
        self._put(par._p)
        cap = self._clone("figcap")
        proto_runs = self.proto["figcap"].findall(qn("w:r"))
        cap.append(self._run("FIGURE %d.\u2003" % number, proto_runs[0]))
        text = path.stem[7:] + (". " + self.captions[number] if number in self.captions else "")
        cap.append(self._run(text, proto_runs[-1]))
        ordered_insert(cap.find(qn("w:pPr")), self._jc("center"), PPR_SEQ)
        self._bookmark(cap, "_AdaFig%d" % number)
        self._put(cap)
        self.figures.append(number)
        self.first_par = True

    @staticmethod
    def _jc(val):
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), val)
        return jc

    # ---------- section files ----------
    def section_file(self, fn):
        text = read_section(fn)
        tbl, table_cap, sub = [], None, 0
        is_refs = fn.startswith("24-")
        for ln in text.splitlines() + [""]:
            if ln.startswith("|"):
                tbl.append(ln)
                continue
            if tbl:
                self.table(tbl, *(table_cap or ()))
                tbl, table_cap = [], None
            m = TABLE_CAPTION.match(ln)
            if m:
                table_cap = (int(m[1]), m[2].strip())
            elif ln.startswith("# "):
                title = ln[2:].strip()
                num = re.match(r"^(\d+)\.\s+(.*)$", title)
                label = ("%s.\u2003%s" % (ROMAN[int(num[1])], num[2].upper())) if num else title.upper()
                name = "_AdaSec%d" % (len(self.sections) + 1)
                self.sections.append({"label": label.replace("\u2003", " "), "title": num[2] if num else title,
                                      "numbered": bool(num), "bookmark": name,
                                      "words": None if is_refs else n_words(text)})
                self.heading(1, label, name)
                sub = 0
            elif ln.startswith("## "):
                head = re.sub(r"^(\d+\.\d+|[A-Z]\.)\s+", "", ln[3:].strip())
                self.heading(2, "%s.\u2002%s" % (chr(ord("A") + sub), head.upper()))
                sub += 1
            elif ln.startswith(("- ", "* ")):
                self.paragraph("•\u00a0" + ln[2:])
            elif ln.strip() and not self.equation(ln):
                self.paragraph(ln.strip(), size=16 if is_refs else None, hanging=is_refs)

    # ---------- front matter ----------
    def front_matter(self, abstract, exec_lines, keywords):
        paras = self.doc.paragraphs
        for r in paras[P_COVER_TITLE].runs:                     # cover: On "<Title>"
            if r.text in ("<", ">"):
                r.text = ""
            elif "Title" in r.text:
                r.text = r.text.replace("Title", TITLE)
            sz = r._r.find(qn("w:rPr") + "/" + qn("w:sz"))
            if sz is not None and sz.get(qn("w:val")) == "64":
                for tag in ("w:sz", "w:szCs"):
                    el = r._r.find(qn("w:rPr") + "/" + qn(tag))
                    if el is not None:
                        el.set(qn("w:val"), COVER_TITLE_SZ)
        for t in self.fm[P_TITLE].iter(qn("w:t")):              # template: title in title case
            if "TITLE OF THE PAPER" in (t.text or ""):
                t.text = TITLE
        self._authors()
        self._team_id()
        # ABSTRACT: one self-contained paragraph
        abs_p = self.fm[P_ABSTRACT]
        runs = abs_p.findall(qn("w:r"))
        for r in runs[1:]:
            abs_p.remove(r)
        self._inline(abs_p, " " + abstract, base_run=runs[1] if len(runs) > 1 else None)
        self._bookmark(abs_p, "_AdaSec0")
        terms_p = self.fm[P_TERMS]
        tr = terms_p.findall(qn("w:r"))
        for r in tr[1:]:
            terms_p.remove(r)
        terms_p.append(self._run(" " + ", ".join(sorted(keywords, key=str.lower)) + ".",
                                 tr[1] if len(tr) > 1 else None))
        # EXECUTIVE SUMMARY with Table 1, in the single-column opening section after INDEX TERMS
        saved, self.anchor = self.anchor, self.fm[P_TERMS + 1]
        self.heading(1, "EXECUTIVE SUMMARY", "_AdaSecES")
        tbl, cap = [], None
        for ln in exec_lines[1:] + [""]:
            if ln.startswith("|"):
                tbl.append(ln)
                continue
            if tbl:
                self.in_block = True                           # section 0 is single-column already
                self.table(tbl, *cap)
                self.in_block = False
                tbl = []
            m = TABLE_CAPTION.match(ln)
            if m:
                cap = (int(m[1]), m[2].strip())
            elif ln.strip():
                self.paragraph(ln.strip())
        self.anchor = saved

    def _authors(self):
        """Author and e-mail lines: keep the guide (template author 1), flag the student slots."""
        auth = self.fm[P_AUTHORS]
        runs = auth.findall(qn("w:r"))
        name_proto, sup_proto = runs[1], runs[2]               # "M Rajasekhara Babu", "1"
        for r in runs[3:]:
            auth.remove(r)
        for i in range(1, STUDENTS + 1):
            auth.append(self._run(", ", name_proto))
            auth.append(self._run("[Student %d name — to be entered]" % i, name_proto, flag=True))
            auth.append(self._run(str(i + 1), sup_proto))
        mail = self.fm[P_EMAILS]
        tail = mail.findall(qn("w:r"))                         # runs after the guide's mailto link
        text_proto, sup_proto = copy.deepcopy(tail[0]), copy.deepcopy(tail[1])
        for r in tail:
            mail.remove(r)
        for i in range(1, STUDENTS + 1):
            mail.append(self._run(", ", text_proto))
            mail.append(self._run(str(i + 1), sup_proto))
            mail.append(self._run("[Student %d VIT e-mail — to be entered]" % i, text_proto, flag=True))

    def _team_id(self):
        """Cover text box (DrawingML and its VML fallback): flag the empty Team ID field."""
        for t in list(self.doc.element.body.iter(qn("w:t"))):
            if (t.text or "").strip() == "Team ID:" and any(
                    a.tag == qn("w:txbxContent") for a in t.iterancestors()):
                run = t.getparent()
                run.addnext(self._run("[to be entered]", run))

    # ---------- front-matter tables ----------
    @staticmethod
    def _cell_rpr(p):
        r = p.find(qn("w:r"))
        if r is not None and r.find(qn("w:rPr")) is not None:
            return copy.deepcopy(r.find(qn("w:rPr")))
        rpr = OxmlElement("w:rPr")
        mark = p.find(qn("w:pPr") + "/" + qn("w:rPr"))
        for c in (mark if mark is not None else ()):
            if c.tag.split("}")[1] in RPR_ALLOWED:
                rpr.append(copy.deepcopy(c))
        return rpr

    def _cell(self, tc, text, flag=False):
        """Write text into a template cell, keeping its paragraph (numbering, spacing) and font."""
        ps = tc.findall(qn("w:p"))
        for extra in ps[1:]:
            tc.remove(extra)
        p = ps[0]
        rpr = self._cell_rpr(p)
        for child in list(p):
            if child.tag != qn("w:pPr"):
                p.remove(child)
        if text:
            if flag:
                highlight(rpr)
            r = OxmlElement("w:r")
            r.append(rpr)
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = text
            r.append(t)
            p.append(r)
        return p

    def _pageref(self, tc, bookmark):
        p = self._cell(tc, "")
        rpr = self._cell_rpr(p)
        for kind, text in (("begin", None), ("instr", " PAGEREF %s \\h " % bookmark),
                           ("separate", None), ("text", "–"), ("end", None)):
            r = OxmlElement("w:r")
            r.append(copy.deepcopy(rpr))
            if kind == "instr":
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            elif kind == "text":
                el = OxmlElement("w:t")
                el.text = text
                self.pagerefs.append((bookmark, el))
            else:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
            r.append(el)
            p.append(r)

    @staticmethod
    def _rows(table):
        return table._tbl.findall(qn("w:tr"))

    @staticmethod
    def _align_like_first_rows(table):
        """The template's own empty cells switch to right/bottom alignment from about row 8 on,
        which only shows once text is filled in; align every row like the first ones."""
        for tc in table._tbl.iter(qn("w:tc")):
            va = tc.find(qn("w:tcPr") + "/" + qn("w:vAlign"))
            if va is not None:
                va.getparent().remove(va)
            for jc in tc.iter(qn("w:jc")):
                if jc.get(qn("w:val")) in ("right", "end"):
                    jc.getparent().remove(jc)

    @staticmethod
    def _tcs(tr):
        return tr.findall(qn("w:tc"))

    def fill_tables(self, counts):
        """Index, List of Figures, List of Tables, Contributions and the consolidated report.

        Template layout kept: unnumbered rows for front/back matter, the template's auto-numbered
        rows 1.-22. for the 22 sections, its fonts, spacing and banding."""
        t_index, t_figs, t_tabs, t_contrib, t_total = self.doc.tables[:5]
        numbered = [s for s in self.sections if s["numbered"]]
        back = [s for s in self.sections if not s["numbered"]]          # DA, References, Appendix
        assert len(numbered) == 22 and [s["title"] for s in back] == \
            ["Data Availability", "References", "Appendix"], [s["title"] for s in back]

        def layout(table, front_labels, n_cols):
            rows = self._rows(table)
            first, nums, refs, app = rows[1], rows[2:24], rows[24], rows[25]
            for tr in rows[26:]:
                table._tbl.remove(tr)
            fronts = [first]
            for _ in front_labels[1:]:
                tr = copy.deepcopy(first)
                fronts[-1].addnext(tr)
                fronts.append(tr)
            da = copy.deepcopy(refs)
            refs.addprevious(da)
            return list(zip(fronts, front_labels)) + list(zip(nums, numbered)) + \
                list(zip((da, refs, app), back))

        index_front = [("Abstract", "_AdaSec0"), ("Executive Summary", "_AdaSecES")]
        for tr, item in layout(t_index, index_front, 3):
            label, bm = item if isinstance(item, tuple) else (item["title"], item["bookmark"])
            tcs = self._tcs(tr)
            self._cell(tcs[1], label)
            self._pageref(tcs[2], bm)

        contrib_front = [("Abstract", counts["abstract"]), ("Index Terms (Keywords)", counts["keywords"]),
                         ("Executive Summary (with Table 1)", counts["exec"])]
        for tr, item in layout(t_contrib, contrib_front, 4):
            if isinstance(item, tuple):
                label, words = item
            else:
                label, words = item["title"], item["words"]
            tcs = self._tcs(tr)
            self._cell(tcs[1], label)
            self._cell(tcs[2], "[to be entered]", flag=True)
            self._cell(tcs[3], "Excl." if words is None else str(words))

        rows = self._rows(t_total)
        for tr in rows[2:2 + STUDENTS]:
            tcs = self._tcs(tr)
            self._cell(tcs[1], "[to be entered]", flag=True)
            self._cell(tcs[2], "[to be entered]", flag=True)
        self._cell(self._tcs(rows[-1])[-1], str(counts["total"]))

        names = {int(p.name[4:6]): p.stem[7:] for p in DIAGRAMS.glob("Fig *.png")}
        rows = self._rows(t_figs)
        while len(rows) - 1 < len(self.figures):       # the copied rows continue the auto-numbering
            rows[-1].addnext(copy.deepcopy(rows[-1]))
            rows = self._rows(t_figs)
        for tr in rows[1 + len(self.figures):]:
            t_figs._tbl.remove(tr)
        for tr, n in zip(self._rows(t_figs)[1:], sorted(self.figures)):
            tcs = self._tcs(tr)
            self._cell(tcs[1], names[n])
            self._pageref(tcs[2], "_AdaFig%d" % n)
        rows = self._rows(t_tabs)
        for tr in rows[1 + len(self.tables):]:
            t_tabs._tbl.remove(tr)
        for tr, (n, title) in zip(self._rows(t_tabs)[1:], sorted(self.table_titles.items())):
            tcs = self._tcs(tr)
            self._cell(tcs[1], title)
            self._pageref(tcs[2], "_AdaTab%d" % n)
        for table in (t_index, t_figs, t_tabs, t_contrib):
            self._align_like_first_rows(table)

    def paginate_front_matter(self):
        for i in SPACERS:
            p = self.fm[i]
            assert not "".join(t.text or "" for t in p.iter(qn("w:t"))).strip(), i
            p.getparent().remove(p)
        for i in PAGE_STARTS:
            ppr = self.fm[i].find(qn("w:pPr"))
            if ppr is None:
                ppr = OxmlElement("w:pPr")
                self.fm[i].insert(0, ppr)
            ordered_insert(ppr, OxmlElement("w:pageBreakBefore"), PPR_SEQ)

    # ---------- package parts ----------
    def _part(self, name):
        for part in self.doc.part.package.iter_parts():
            if str(part.partname) == name:
                return part
        raise KeyError(name)

    def footers_and_cover(self):
        # template footer prints a literal "1" on every page: make it a PAGE field in all footers
        proto = self._part("/word/footer1.xml").element
        para = copy.deepcopy(proto.find(qn("w:p")))
        run = para.find(qn("w:r"))
        rpr = run.find(qn("w:rPr"))
        para.remove(run)
        tab = OxmlElement("w:r")
        tab.append(copy.deepcopy(rpr))
        tab.append(OxmlElement("w:tab"))
        para.append(tab)
        for kind, text in (("begin", None), ("instr", " PAGE "), ("separate", None), ("t", "1"),
                           ("end", None)):
            r = OxmlElement("w:r")
            r.append(copy.deepcopy(rpr))
            if kind == "instr":
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            elif kind == "t":
                el = OxmlElement("w:t")
                el.text = text
            else:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
            r.append(el)
            para.append(r)
        for name in ("/word/footer1.xml", "/word/footer2.xml", "/word/footer3.xml"):
            ftr = self._part(name).element
            for p in ftr.findall(qn("w:p")):
                ftr.remove(p)
            ftr.append(copy.deepcopy(para))
        # cover SmartArt (data model and its cached drawing): flag the student banners
        for name in ("/word/diagrams/data1.xml", "/word/diagrams/drawing1.xml"):
            part = self._part(name)
            blob = part.blob.decode("utf-8")
            blob, n_reg = re.subn(r"<a:t>Reg No ?</a:t>", "<a:t>Reg No: [to be entered]</a:t>", blob)
            blob, n_name = re.subn(r"<a:t>Name</a:t>", "<a:t>Name: [to be entered]</a:t>", blob)
            assert n_reg == n_name == STUDENTS, (name, n_reg, n_name)
            part._blob = blob.encode("utf-8")

    def properties(self, keywords, stats=None):
        cp = self.doc.core_properties
        cp.title, cp.keywords = TITLE, "; ".join(sorted(keywords, key=str.lower))
        cp.author, cp.last_modified_by, cp.revision = "", "", 1
        cp.created = cp.modified = datetime.now(timezone.utc).replace(microsecond=0)
        app = self._part("/docProps/app.xml")
        xml = app.blob.decode("utf-8")
        for tag in ("Characters", "CharactersWithSpaces", "Lines", "Paragraphs"):
            xml = re.sub(r"<%s>[^<]*</%s>" % (tag, tag), "", xml)
        xml = re.sub(r"<TotalTime>[^<]*</TotalTime>", "<TotalTime>0</TotalTime>", xml)
        if stats:
            xml = re.sub(r"<Pages>[^<]*</Pages>", "<Pages>%d</Pages>" % stats[0], xml)
            xml = re.sub(r"<Words>[^<]*</Words>", "<Words>%d</Words>" % stats[1], xml)
        app._blob = xml.encode("utf-8")

    # ---------- page numbers ----------
    def word_page_numbers(self):
        """Let Word paginate a temporary copy; return ({bookmark: page}, (pages, words)) or None."""
        with tempfile.TemporaryDirectory() as tmp:
            src, dst = Path(tmp, "provisional.docx"), Path(tmp, "updated.docx")
            self.doc.save(src)
            script = PS_WORD_UPDATE % {"src": src, "dst": dst}
            try:
                res = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
                                     capture_output=True, text=True, timeout=900)
            except (OSError, subprocess.TimeoutExpired) as exc:
                print("Word step unavailable:", exc)
                return None
            m = re.search(r"STATS (\d+) (\d+)", res.stdout)
            if res.returncode or not m or not dst.exists():
                print("Word step failed:", res.stdout[-400:], res.stderr[-400:])
                return None
            with zipfile.ZipFile(dst) as z:
                root = etree.fromstring(z.read("word/document.xml"))
        pages, stack = {}, []
        for el in root.iter(qn("w:fldChar"), qn("w:instrText"), qn("w:t")):
            if el.tag == qn("w:fldChar"):
                kind = el.get(qn("w:fldCharType"))
                if kind == "begin":
                    stack.append({"instr": "", "result": "", "in_result": False})
                elif kind == "separate" and stack:
                    stack[-1]["in_result"] = True
                elif kind == "end" and stack:
                    f = stack.pop()
                    m2 = re.match(r"\s*PAGEREF\s+(\S+)", f["instr"])
                    if m2:
                        pages[m2[1]] = f["result"].strip()
            elif stack and el.tag == qn("w:instrText") and not stack[-1]["in_result"]:
                stack[-1]["instr"] += el.text or ""
            elif stack and el.tag == qn("w:t") and stack[-1]["in_result"]:
                stack[-1]["result"] += el.text or ""
        return pages, (int(m[1]), int(m[2]))


def main():
    b = Builder()
    b.table_titles = {}
    for fn in ["00-exec-summary.md"] + BODY:
        for ln in read_section(fn).splitlines():
            m = TABLE_CAPTION.match(ln)
            if m:
                b.table_titles[int(m[1])] = m[2].strip()
    abstract_lines = [l.strip() for l in read_section("00-abstract.md").splitlines()
                      if l.strip() and not l.startswith("#")]
    exec_text = read_section("00-exec-summary.md")
    keywords = [k.strip() for k in read_section("01-keywords.md").splitlines()[-1].split(",") if k.strip()]
    b.front_matter(abstract_lines[0], exec_text.splitlines(), keywords)
    figs = {int(p.name[4:6]): p for p in DIAGRAMS.glob("Fig *.png")}
    for fn in BODY:
        b.section_file(fn)
        for n in sorted(k for k, v in FIG_AFTER.items() if fn.startswith(v)):
            if n in figs:
                b.figure(n, figs[n])
    b._close_block()

    counts = {"abstract": n_words(abstract_lines[0]),
              "keywords": n_words(read_section("01-keywords.md").splitlines()[-1]),
              "exec": n_words(exec_text)}
    counts["total"] = counts["abstract"] + counts["keywords"] + counts["exec"] + \
        sum(s["words"] for s in b.sections if s["words"] is not None)
    b.fill_tables(counts)
    b.paginate_front_matter()
    b.footers_and_cover()
    b.properties(keywords)

    result = b.word_page_numbers()
    settings = b.doc.settings.element
    if result:
        pages, stats = result
        missing = sorted({bm for bm, _t in b.pagerefs} - set(pages))
        if missing or not all(re.fullmatch(r"\d+", pages[bm]) for bm, _t in b.pagerefs):
            sys.exit("BUILD CHECK FAILED: page numbers not resolved for %s" % (missing or pages))
        for bm, t in b.pagerefs:
            t.text = pages[bm]
        b.properties(keywords, stats)
    else:
        upd = OxmlElement("w:updateFields")                 # let Word fill them on open
        upd.set(qn("w:val"), "true")
        settings.append(upd)
    b.doc.save(OUT)

    problems = []
    if sorted(b.figures) != list(range(1, 33)):
        problems.append(f"figures placed {sorted(b.figures)}")
    if b.equations != list(range(1, 14)):
        problems.append(f"equations not displayed in order 1-13: {b.equations}")
    if sorted(b.tables) != list(range(1, 16)):
        problems.append(f"tables {sorted(b.tables)}")
    n_omml = len(etree.XPath(".//m:oMath", namespaces={"m": M_NS})(b.doc.element))
    if not 150 <= counts["abstract"] <= 250 or len(abstract_lines) != 1:
        problems.append(f"abstract: {len(abstract_lines)} paragraph(s), {counts['abstract']} words")
    if not WORD_BAND[0] <= counts["total"] <= WORD_BAND[1]:
        problems.append(f"total words {counts['total']} outside {WORD_BAND}")

    def is_code(run):
        fonts = run.find(qn("w:rPr") + "/" + qn("w:rFonts"))
        return fonts is not None and fonts.get(qn("w:ascii")) == "Consolas"
    joined = "".join(t.text or "" for r in b.doc.element.iter(qn("w:r")) if not is_code(r)
                     for t in r.iter(qn("w:t")))
    for leak in ("**", "[@", "<TITLE", "<Title>", "TITLE OF THE PAPER", "Second B. Author",
                 "Third C. Author", "Fourth D. Author", "Fifth E. Author", "Vit email"):
        if leak in joined:
            problems.append(f"leak {leak!r}")
    body_text = "\n".join(read_section(fn) for fn in ["00-exec-summary.md"] + BODY
                          if not fn.startswith("24-"))
    body_text = figure_refs(body_text)
    explicit = {int(n) for n in re.findall(r"\bFig\. (\d+)", body_text)}
    explicit |= {int(n) for pair in re.findall(r"\bFigs\. (\d+) and (\d+)", body_text) for n in pair}
    if set(range(1, 33)) - explicit:
        problems.append(f"figures without an explicit text reference: {sorted(set(range(1, 33)) - explicit)}")
    tabs = {int(n) for n in re.findall(r"\bTables? (\d+)", body_text)}
    if set(range(1, 16)) - tabs:
        problems.append(f"tables never referenced: {sorted(set(range(1, 16)) - tabs)}")
    prose = "\n".join(l for l in body_text.splitlines() if not EQUATION_LINE.match(l.strip())
                      or not re.search(r"[=←]", l))
    eqs = {int(n) for n in re.findall(r"\((\d{1,2})\)", prose)}
    if set(range(1, 14)) - eqs:
        problems.append(f"equations never referenced: {sorted(set(range(1, 14)) - eqs)}")
    if re.search(r"(?<!PLAN )\bSections? \d", body_text):
        problems.append("Arabic section cross-reference left: %s" % re.findall(r".{20}\bSections? \d.{10}", body_text)[:3])
    cites = {int(n) for n in re.findall(r"\[(\d{1,3})\]", body_text)}
    refs = {int(n) for n in re.findall(r"^\[(\d+)\]", read_section("24-references.md"), re.M)}
    if cites - refs or refs - cites:
        problems.append(f"citations without reference {sorted(cites - refs)}; uncited {sorted(refs - cites)}")
    print("wrote %s: %d figures, %d tables, %d OMML equations, %d sections, %d words excl. references"
          " (abstract %d, executive summary %d), page numbers %s"
          % (OUT.name, len(b.figures), len(b.tables), n_omml, len(b.sections), counts["total"],
             counts["abstract"], counts["exec"], "from Word" if result else "left for Word to update"))
    if problems:
        sys.exit("BUILD CHECK FAILED: " + "; ".join(problems))


if __name__ == "__main__":
    main()
